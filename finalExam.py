# ALEDUK5688
import csv
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from datetime import date

path = "C:\\FinalExam\\final.csv"

def askUser():
    totalSum = 0
    # For loop will loop 5 times while asking for a number. Every loop will
    # add the number input to totalSum. After the loop is over, it'll
    # output the totalSum variable.
    for i in range(5):
        num = int(input("Please enter a number: "))
        totalSum += num
    print(f'The total of the five numbers is {totalSum}.')


def askIncome():
    with open(path, 'a') as f:
        # For loop will loop 5 times asking for a name and a paired income amount.
        # Every entry will be appended to the file in 'path'
        for i in range(5):
            name = input("Enter the person's name: ")
            income = input("Enter the person's income: ")

            f.write(name + ',' + income + "\n")     


def excelPie():
    workbook = Workbook()
    sheet = workbook.active

    with open(path, 'r')as f:
        reader = csv.reader(f)

        # For loop reads out each row from file in 'path' before placing the
        # people's names and incomes into the .xlsx file
        row = 1
        for data in reader:
            sheet.cell(row=row, column=1, value=data[0])
            sheet.cell(row=row, column=2, value=int(data[1]))
            row += 1
        # This creates a new pie chart
        chart = PieChart()
        # Selects the names from column A to use as labels for chart
        labels = Reference(sheet, min_col=1, min_row=1, max_row=row-1)
        # Selects the incomes from column B to use as values for the chart
        values = Reference(sheet, min_col=2, min_row=1, max_row=row-1)
        # Adds the income values to the chart
        chart.add_data(values, titles_from_data=False)
        # Assigns names as labels for the pie chart slices
        chart.set_categories(labels)
        # Creates the title for the chart using my Student ID and the date
        chart.title = "ALEDUK5688 " + date.today().strftime("%B %d, %Y")
        # Places the pie chart onto the worksheet
        sheet.add_chart(chart, "D2")
        # saves the worksheet as the file labeled below
        workbook.save("final.xlsx")

def verticalBar():
    names = []
    incomes = []

    with open(path, 'r') as f:
        reader = csv.reader(f)

        # For loop reads each row from 'path' and stores the values in separate lists
        # in order to make the bar graph
        for data in reader:
            names.append(data[0])
            incomes.append(int(data[1]))

        plt.bar(names, incomes)
        plt.title("ALEDUK5688 " + date.today().strftime("%B %d, %Y"))
        plt.xlabel("Name")
        plt.ylabel("Annual Income")

        plt.show()
        
def main():
    askUser()
    askIncome()
    excelPie()
    verticalBar()

main()
